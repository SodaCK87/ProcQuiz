"""轉檔正確度的測試。執行：python -m unittest discover -s tests -v

破壞型測試都做在暫存複本上，data/source/ 的原檔一個字都不碰。
"""
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import convert  # noqa: E402

TF = ROOT / "data" / "source" / "true-false.xlsx"
MC = ROOT / "data" / "source" / "multiple-choice.xlsx"


class PureFunctions(unittest.TestCase):
    def test_split_options(self):
        stem, opts = convert.split_options("何者正確？(1)甲。(2)乙。（3）丙。(４)丁。", "t")
        self.assertEqual(stem, "何者正確？")
        self.assertEqual(opts, ["甲。", "乙。", "丙。", "丁。"])

    def test_split_options_rejects_missing_option(self):
        with self.assertRaises(convert.BankError):
            convert.split_options("何者？(1)甲(2)乙(4)丁", "t")

    def test_parse_answer(self):
        self.assertEqual(convert.parse_answer("Ｘ", "true-false", "t"), "X")
        self.assertEqual(convert.parse_answer(" o ", "true-false", "t"), "O")
        self.assertEqual(convert.parse_answer("3", "multiple-choice", "t"), 3)
        for bad, kind in (("Y", "true-false"), ("5", "multiple-choice"), (None, "true-false")):
            with self.assertRaises(convert.BankError):
                convert.parse_answer(bad, kind, "t")


class Committed(unittest.TestCase):
    def test_committed_json_matches_regeneration(self):
        # 產出物進版控，這條守「改了 xlsx 或腳本卻忘了重跑」
        for kind in convert.KINDS:
            data, _ = convert.convert(kind, convert.SOURCE / f"{kind}.xlsx")
            committed = (convert.OUT / f"{kind}.json").read_text(encoding="utf-8")
            self.assertEqual(committed, convert.render(data), kind)


class Corruption(unittest.TestCase):
    """陽性對照：每一種題庫錯誤都要讓轉檔中止。"""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _copy(self, src, mutate):
        dst = self.tmp / src.name
        wb = openpyxl.load_workbook(src)
        mutate(wb)
        wb.save(dst)
        return dst

    @staticmethod
    def _row(ws, no):
        for row in ws.iter_rows(min_row=3):
            if str(row[0].value or "").strip() == str(no):
                return row
        raise AssertionError(f"找不到第 {no} 題")

    def test_roundtrip_untouched_copy_is_identical(self):
        # 陰性對照：原樣存回不改任何東西，結果要與原檔一致，後面的紅才歸得到破壞頭上
        for src, kind in ((TF, "true-false"), (MC, "multiple-choice")):
            got, _ = convert.convert(kind, self._copy(src, lambda wb: None))
            want, _ = convert.convert(kind, src)
            self.assertEqual(got, want, kind)

    def _assert_rejects(self, src, kind, mutate, expect):
        path = self._copy(src, mutate)
        with self.assertRaises(convert.BankError) as cm:
            convert.convert(kind, path)
        self.assertIn(expect, str(cm.exception))

    def test_flipped_tf_answer_breaks_o_count(self):
        def flip(wb):
            row = self._row(wb["1-政府採購全生命週期概論"], 1)
            row[2].value = "X" if row[2].value == "O" else "O"
        self._assert_rejects(TF, "true-false", flip, "課程 1")

    def test_known_mismatch_does_not_mask_further_change(self):
        # PQZ-01 的放行只認 76／77；課程 5 再多錯一題就要中止
        def flip(wb):
            ws = wb["5-工程及技術服務採購作業"]
            row = next(r for r in ws.iter_rows(min_row=3) if r[2].value == "X")
            row[2].value = "O"
        self._assert_rejects(TF, "true-false", flip, "課程 5")

    def test_deleted_question_breaks_count(self):
        def delete(wb):
            ws = wb["3-政府採購法之履約管理及驗收"]
            ws.delete_rows(self._row(ws, 97)[0].row)
        self._assert_rejects(TF, "true-false", delete, "目錄頁寫 97 題")

    def test_invalid_answer(self):
        def bad(wb):
            self._row(wb["1-政府採購全生命週期概論"], 2)[2].value = "5"
        self._assert_rejects(MC, "multiple-choice", bad, "選擇題答案應為 1–4")

    def test_missing_option(self):
        def bad(wb):
            row = self._row(wb["1-政府採購全生命週期概論"], 1)
            row[3].value = row[3].value.replace("(3)", "")
        self._assert_rejects(MC, "multiple-choice", bad, "選項標記")

    def test_duplicate_with_conflicting_answer(self):
        def bad(wb):
            row = self._row(wb["9-錯誤採購態樣"], 78)  # 與 mc-06-0109 同題
            row[2].value = (int(row[2].value) % 4) + 1
        self._assert_rejects(MC, "multiple-choice", bad, "答案不一致")


if __name__ == "__main__":
    unittest.main()
