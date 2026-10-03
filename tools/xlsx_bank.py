"""讀第三方整理的「題庫解析」xlsx，只取它的「依據法源」「解析」兩欄來補官方題庫。

題目與答案以官方題庫為準（見 official.py）；這裡仍逐課程核對目錄頁，
因為解析掛錯題比沒有解析更糟，來源本身不一致時要先停下來。
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

KINDS = {
    "true-false": {"label": "是非題", "prefix": "tf"},
    "multiple-choice": {"label": "選擇題", "prefix": "mc"},
}

OPTION_MARK = re.compile(r"[(（]([1-4１-４])[)）]")

# 題目列照固定欄位位置讀（read_xlsx），欄一少或對調不會報錯、只會把法源當解析掛上去，
# 所以先核對第 1 列標題；只核程式真的會讀的那幾欄（練習欄不讀）。
HEADERS = {0: "編號", 2: "答案", 3: "試題(點擊可回目錄)", 4: "依據法源", 5: "解析"}

# 已查過、確認是 xlsx 本身的不一致（見 docs/問題台帳.md），只在數字完全相同時放行。
# 換版後數字變了就照樣中止，逼人重新查一次。
KNOWN_TOC_O_MISMATCH = {
    ("115.7.23", 5): (76, 77),  # PQZ-01
}


class BankError(Exception):
    pass


def text_of(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def norm_header(v) -> str:
    # 全半形與空白不計，但不做包含比對：「解析」與「解析（補充）」要算不同
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(v or "")))


def check_headers(ws, kind: str, path: Path) -> None:
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for idx, want in HEADERS.items():
        got = header[idx] if idx < len(header) else None
        if norm_header(got) != norm_header(want):
            raise BankError(f"xlsx {KINDS[kind]['label']} {path.name} 工作表「{ws.title}」第 1 列 {get_column_letter(idx + 1)} 欄："
                            f"標題應為「{want}」，實際「{got if got is not None else ''}」")


def read_toc(ws, kind: str) -> tuple[str, dict[int, dict]]:
    """回傳（版次, {課程編號: {name, count, o_count}}）。目錄頁左右兩欄並排。"""
    version = None
    for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
        for v in row:
            m = re.search(r"(\d+\.\d+\.\d+)版", str(v or ""))
            if m:
                version = m.group(1)
    if not version:
        raise BankError(f"{kind}：目錄頁找不到版次字樣（例如 115.7.23版）")

    width = 6 if kind == "true-false" else 4  # 是非題多「答案題數 O／X」兩欄
    toc: dict[int, dict] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        for base in (0, width):
            cells = row[base:base + width]
            if len(cells) < 3 or not isinstance(cells[0], int) or not isinstance(cells[2], int):
                continue
            entry = {"name": str(cells[1]).strip(), "count": cells[2]}
            if kind == "true-false":
                entry["o_count"] = cells[4]
            toc[cells[0]] = entry
    if sorted(toc) != list(range(1, 15)):
        raise BankError(f"{kind}：目錄頁應有課程 1–14，實際讀到 {sorted(toc)}")
    return version, toc


def split_options(text: str, where: str) -> tuple[str, list[str]]:
    marks = list(OPTION_MARK.finditer(text))
    nums = [unicodedata.normalize("NFKC", m.group(1)) for m in marks]
    if nums != ["1", "2", "3", "4"]:
        raise BankError(f"{where}：選項標記應為 (1)(2)(3)(4)，實際 {nums}")
    stem = text[:marks[0].start()].strip()
    options = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        options.append(text[m.end():end].strip())
    if not stem or any(not o for o in options):
        raise BankError(f"{where}：題幹或選項為空")
    return stem, options


def parse_answer(raw, kind: str, where: str):
    s = unicodedata.normalize("NFKC", str(raw or "")).strip().upper()
    if kind == "true-false":
        if s not in ("O", "X"):
            raise BankError(f"{where}：是非題答案應為 O／X，實際 {raw!r}")
        return s
    if s not in ("1", "2", "3", "4"):
        raise BankError(f"{where}：選擇題答案應為 1–4，實際 {raw!r}")
    return int(s)


def read_xlsx(kind: str, path: Path) -> tuple[dict, list[str]]:
    """回傳（{version, courses, questions}, 提醒清單）。questions 的 text 是題目原文（選擇題含選項）。"""
    wb = openpyxl.load_workbook(path)
    version, toc = read_toc(wb.worksheets[0], kind)
    label = KINDS[kind]["label"]
    notes: list[str] = []
    questions = []

    for ws in wb.worksheets[1:]:
        m = re.match(r"(\d+)-(.+)", ws.title)
        if not m:
            raise BankError(f"{kind}：工作表名稱不是「編號-課程」格式：{ws.title}")
        course = int(m.group(1))
        if course not in toc or toc[course]["name"] != m.group(2).strip():
            raise BankError(f"{kind}：工作表「{ws.title}」與目錄頁課程名稱對不上")
        check_headers(ws, kind, path)

        seen: set[int] = set()
        for row in ws.iter_rows(min_row=3):
            no_text = text_of(row[0].value)
            if no_text is None or not no_text.isdigit():
                # 編號欄非數字的列是課程說明或分組標題，不是題目
                if no_text and any(text_of(c.value) for c in row[2:4]):
                    raise BankError(f"{ws.title} 第 {row[0].row} 列：編號不是數字卻有答案或題目")
                continue
            no = int(no_text)
            where = f"xlsx {label} {ws.title} 第 {no} 題"
            if no in seen:
                raise BankError(f"{where}：編號重複")
            seen.add(no)

            answer = parse_answer(row[2].value, kind, where)
            text = text_of(row[3].value)
            if not text:
                raise BankError(f"{where}：題目為空")
            if kind == "multiple-choice":
                split_options(text, where)
            questions.append({"id": f"{KINDS[kind]['prefix']}-{course:02d}-{no:04d}", "course": course,
                              "no": no, "answer": answer, "text": text,
                              "law": text_of(row[4].value), "explanation": text_of(row[5].value)})

        if sorted(seen) != list(range(1, len(seen) + 1)):
            raise BankError(f"{kind} {ws.title}：編號不是 1–{len(seen)} 連號")

    for course, entry in toc.items():
        qs = [q for q in questions if q["course"] == course]
        if len(qs) != entry["count"]:
            raise BankError(f"xlsx {kind} 課程 {course}：目錄頁寫 {entry['count']} 題，實際 {len(qs)} 題")
        if kind == "true-false":
            o = sum(q["answer"] == "O" for q in qs)
            if KNOWN_TOC_O_MISMATCH.get((version, course)) == (entry["o_count"], o):
                notes.append(f"xlsx 是非題 課程 {course}：目錄頁寫答案 O 有 {entry['o_count']} 題，實際 {o} 題（已知，PQZ-01）")
            elif o != entry["o_count"]:
                raise BankError(f"xlsx {kind} 課程 {course}：目錄頁寫答案 O 有 {entry['o_count']} 題，實際 {o} 題")

    return {"version": version,
            "courses": {c: toc[c]["name"] for c in sorted(toc)},
            "questions": questions}, notes
